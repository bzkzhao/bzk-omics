"""Round-15b descriptive, UNREGISTERED: the protein-grain census, with the mapping tested first.

Usage, from the repository root:
    .venv/bin/python notes/scripts/diagnose_round15b.py 2>&1 | tee notes/logs/round15b.txt

Round 15's JJ was void: the deposited `proteinGroups` names its samples `KO_INF_P_2hGradient1`
while MOESM4 names them `KO_IFN_1`, so the naive mapping matched nothing and every cell came back
"drawn". No committed curation record maps the protein columns.

The mapping below is the reviewer's inference, NOT a curated record: INF for IFN, P for proteome,
the trailing digit as the replicate. It is therefore TESTED before it is used, the same way
Data Table S1 was tested against the site table: for cells the deposit provides, the published
value should equal log2 of the deposited one to within rounding. If that fails, the mapping is
wrong and the census is not run.

The question: same samples, same lab, same software, no PTM enrichment — what share of the
published PROTEIN claims contain a value the authors drew? Site-level is 95%.
"""

import math
import warnings

import numpy as np
from openpyxl import load_workbook

from bzk.adapters import maxquant
from bzk.provenance.raw_store import raw_root
from bzk.sources.pxd018299_h10 import HOME

warnings.filterwarnings("ignore")

MAPPING = {
    "LFQ intensity KO_IFN_1": "LFQ intensity KO_INF_P_2hGradient1",
    "LFQ intensity KO_IFN_2": "LFQ intensity KO_INF_P_2hGradient2",
    "LFQ intensity KO_IFN_3": "LFQ intensity KO_INF_P_2hGradient3",
    "LFQ intensity WT_IFN_1": "LFQ intensity WT_INF_P_2hGradient1",
    "LFQ intensity WT_IFN_2": "LFQ intensity WT_INF_P_2hGradient2",
    "LFQ intensity WT_IFN_3": "LFQ intensity WT_INF_P_2hGradient3",
}

moesm4 = next(iter(raw_root(HOME).glob("*/41416_2020_1167_MOESM4_ESM.xlsx")), None)
groups_path = next(iter(raw_root(HOME).glob("*/HAP1_USP18KO_proteinGroups.txt")), None)
if not moesm4 or not groups_path:
    raise SystemExit("MOESM4 or the deposited proteinGroups is not in the raw store")

book = load_workbook(moesm4, read_only=True, data_only=True)
body = [r for r in book.worksheets[0].iter_rows(values_only=True) if any(c is not None for c in r)]
head_index = next(i for i, r in enumerate(body) if any("LFQ intensity" in str(c) for c in r if c))
head = [str(c).strip() if c is not None else "" for c in body[head_index]]
published_rows = body[head_index + 1:]
groups = maxquant.read_table(groups_path)
group_column = {n: i for i, n in enumerate(groups.header)}
lookup: dict[str, list[str]] = {}
for row in groups.rows:
    for part in row[group_column["Protein IDs"]].split(";"):
        lookup.setdefault(part.strip(), row)
print(f"published protein claims: {len(published_rows)}; deposited protein groups: {len(groups.rows)}")


def deposit_value(group: list[str], published_name: str) -> float:
    name = MAPPING[published_name]
    text = group[group_column[name]].strip() if name in group_column else ""
    try:
        return float(text)
    except ValueError:
        return 0.0


pairs, matched = [], []
key_index = head.index("T: Protein IDs")
for row in published_rows:
    key = str(row[key_index]).split(";")[0].strip() if row[key_index] else ""
    group = lookup.get(key)
    if group is None:
        continue
    matched.append((row, group))
    for name in MAPPING:
        published = row[head.index(name)]
        try:
            published = float(published)
        except (TypeError, ValueError):
            continue
        value = deposit_value(group, name)
        if value > 0:
            pairs.append((published, math.log2(value)))

print(f"matched {len(matched)} of {len(published_rows)} published claims to a deposited group")
print("\ntesting the mapping: published against log2 of the deposited value, for cells the deposit provides")
if not pairs:
    raise SystemExit("  no comparable cells — the mapping is wrong and the census is not run")
published_values, deposit_values = np.array(pairs).T
difference = published_values - deposit_values
print(f"  cells compared: {len(pairs)}")
print(f"  |difference| <= 0.01: {int((np.abs(difference) <= 0.01).sum())} ({(np.abs(difference) <= 0.01).mean():.1%})")
print(f"  quantiles of the difference: {[round(float(q), 3) for q in np.percentile(difference, [5, 50, 95])]}")
print(f"  correlation: {np.corrcoef(published_values, deposit_values)[0, 1]:.4f}")
if (np.abs(difference) <= 0.01).mean() < 0.95:
    raise SystemExit("  the mapping does not reproduce the published values: census NOT run")
print("  the mapping is confirmed by the published values themselves")

print("\nthe protein-grain census")
drawn_cells = drawn_claims = whole_arm = 0
for row, group in matched:
    drawn = [deposit_value(group, name) <= 0 for name in MAPPING]
    drawn_cells += sum(drawn)
    drawn_claims += any(drawn)
    whole_arm += all(drawn[:3]) or all(drawn[3:])
print(f"  published protein claims examined: {len(matched)}")
print(f"    containing at least one drawn cell: {drawn_claims} ({drawn_claims / len(matched):.0%})")
print(f"    with a whole arm drawn:             {whole_arm}")
print(f"    drawn cells: {drawn_cells} of {len(matched) * 6} ({drawn_cells / (len(matched) * 6):.1%})")
print("  site level, for comparison: 95% of claims contain a drawn cell, 715 of 798 with a whole arm")
print("\n  Caveat: 25 published protein claims is a small denominator, and MOESM4 is the")
print("  'significantly up' table only, so this is not the whole protein selection.")
