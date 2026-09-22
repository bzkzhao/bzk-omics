"""Round-15 descriptive, UNREGISTERED checks (adversarial review, 2026-09-22).

Usage, from the repository root:
    .venv/bin/python notes/scripts/diagnose_round15.py 2>&1 | tee notes/logs/round15.txt

Not an independent path: the same instruments the registered runs use. Nothing registered changes.

II — THE SIGNED GAP, AND WHAT IMPUTATION PREDICTS. The 43 partially drawn claims are a selected
     class: a claim is in S1 because it was called, and a drawn cell helps it get called when it
     lands low in the wild-type arm. So the absolute gap of 1.055 is upper-leaning. Reported here
     signed, and against two simulations at the recovered parameters (downshift 1.77, width 0.29):
     unconditional, and conditioned on the claim clearing the weakest published statistic. If the
     conditional simulation matches, the distortion is what the authors' own model predicts under
     selection.

JJ — THE PROTEIN-GRAIN CONTROL, INSIDE THE SAME EXPERIMENT. Same samples, same lab, same software,
     no PTM enrichment. MOESM4's 26 published protein claims carry LFQ intensity columns, so the
     census rule applies to them against the deposited `proteinGroups`. This is the first
     within-experiment evidence on the PTM-specific question.

KK — THE 2,341 SCAN, BOUNDED AND GUARDED. The target sits between the localisation-filtered 2,271
     and the contaminant-filtered 2,359, so the paper's filter lies in a small region with two
     monotone axes. Scanned for an EXACT match on distinct modified-peptide ids, reporting how many
     cells hit it and whether any hit is defensible on its face.
"""

import json
import math
import warnings

import numpy as np
from bzk.adapters import maxquant
from bzk.provenance.raw_store import raw_root, verify
from bzk.sources.pxd018299_h10 import (
    HOME,
    PXD018299_SITES,
    SUPP_DATA_1,
    _float_cell,
    _s1_rows,
    _supplement_path,
    anchor_population,
    s1_columns,
)
from openpyxl import load_workbook

warnings.filterwarnings("ignore")
KO = ["Intensity KO_IFN_1", "Intensity KO_IFN_2", "Intensity KO_IFN_3"]
WT = ["Intensity WT_IFN_1", "Intensity WT_IFN_2", "Intensity WT_IFN_3"]
DOWNSHIFT, WIDTH = 1.77, 0.29

table = maxquant.read_table(
    verify(PXD018299_SITES.expected_content_hash, filename=PXD018299_SITES.filename, home=HOME)
)
rows, column = anchor_population(table)
ids = [str(r[column["id"]]).strip() for r in rows]
index_of = {name: i for i, name in enumerate(ids)}


def raw_value(row: list[str], name: str) -> float:
    text = row[column[name]].strip() if name in column else ""
    try:
        return float(text)
    except ValueError:
        return 0.0


def provided(row: list[str], name: str) -> bool:
    return raw_value(row, name) > 0 or any(
        raw_value(row, f"{name}{suffix}") > 0 for suffix in ("___1", "___2", "___3")
    )


with open("tests/fixtures/pxd018299_published_cascade.json") as handle:
    cascade = json.load(handle)["rows"]
header, s1 = _s1_rows(_supplement_path(SUPP_DATA_1, HOME))
published_columns = s1_columns(header)

# Column statistics from the deposit, for the simulation.
column_values = {name: [] for name in KO + WT}
for row in rows:
    for name in KO + WT:
        value = raw_value(row, name)
        if value > 0:
            column_values[name].append(math.log2(value))
column_mean = {name: float(np.mean(v)) for name, v in column_values.items()}
column_sd = {name: float(np.std(v, ddof=1)) for name, v in column_values.items()}

print("II — the signed gap over the 43 partially drawn claims")
signed, records = [], []
for claim in cascade:
    row = rows[index_of[str(claim["deposit_id"])]]
    values = [_float_cell(s1[int(claim["row"])].get(name)) for name in published_columns]
    drawn = [not provided(row, name) for name in KO + WT]
    if not any(drawn) or all(drawn[:3]) or all(drawn[3:]):
        continue
    ko_measured = [v for v, d in zip(values[:3], drawn[:3], strict=True) if not d]
    wt_measured = [v for v, d in zip(values[3:], drawn[3:], strict=True) if not d]
    published_fc = np.mean(values[:3]) - np.mean(values[3:])
    measured_fc = np.mean(ko_measured) - np.mean(wt_measured)
    signed.append(published_fc - measured_fc)
    records.append((str(claim["deposit_id"]), drawn, values, published_fc))
signed = np.array(signed)
print(f"  claims: {len(signed)}")
print(f"  signed (published - measured-only): quartiles {np.percentile(signed, 25):+.3f} / "
      f"{np.median(signed):+.3f} / {np.percentile(signed, 75):+.3f}")
print(f"  published exceeds measured-only in {int((signed > 0).sum())} of {len(signed)} claims")
print(f"  absolute gap (as reported in round 14): median {np.median(np.abs(signed)):.3f}")

print("\n  what the authors' own imputation model predicts for these same claims")
rng = np.random.default_rng(0)
weakest = min(fc for _, _, _, fc in records)
unconditional, conditional = [], []
for _ in range(200):
    for deposit_id, drawn, values, _published in records:
        simulated = list(values)
        for position, name in enumerate(KO + WT):
            if drawn[position]:
                simulated[position] = rng.normal(
                    column_mean[name] - DOWNSHIFT * column_sd[name], WIDTH * column_sd[name]
                )
        ko_measured = [v for v, d in zip(values[:3], drawn[:3], strict=True) if not d]
        wt_measured = [v for v, d in zip(values[3:], drawn[3:], strict=True) if not d]
        simulated_fc = np.mean(simulated[:3]) - np.mean(simulated[3:])
        gap = simulated_fc - (np.mean(ko_measured) - np.mean(wt_measured))
        unconditional.append(gap)
        if simulated_fc >= weakest:
            conditional.append(gap)
print(f"    unconditional: median {np.median(unconditional):+.3f} "
      f"(quartiles {np.percentile(unconditional, 25):+.3f} / {np.percentile(unconditional, 75):+.3f})")
print(f"    conditioned on clearing the weakest published fold change ({weakest:+.3f}): "
      f"median {np.median(conditional):+.3f} over {len(conditional)} of {len(unconditional)} draws")
print(f"    observed: median {np.median(signed):+.3f}")

print("\nJJ — the protein-grain control: MOESM4's published protein claims")
moesm4 = next(iter(raw_root(HOME).glob("*/41416_2020_1167_MOESM4_ESM.xlsx")), None)
groups_path = next(iter(raw_root(HOME).glob("*/HAP1_USP18KO_proteinGroups.txt")), None)
if moesm4 and groups_path:
    book = load_workbook(moesm4, read_only=True, data_only=True)
    sheet = book.worksheets[0]
    body = [r for r in sheet.iter_rows(values_only=True) if any(c is not None for c in r)]
    head_index = next(i for i, r in enumerate(body) if any("LFQ intensity" in str(c) for c in r if c))
    head = [str(c).strip() if c is not None else "" for c in body[head_index]]
    print(f"  MOESM4 header: {head}")
    lfq = [name for name in head if name.startswith("LFQ intensity")]
    groups = maxquant.read_table(groups_path)
    group_column = {n: i for i, n in enumerate(groups.header)}
    deposit_lfq = [n for n in groups.header if n.startswith("LFQ intensity")]
    print(f"  deposited proteinGroups LFQ columns: {deposit_lfq}")
    key_field = next((n for n in head if "protein" in n.lower() or "gene" in n.lower()), None)
    print(f"  MOESM4 identifier column: {key_field!r}")
    if key_field:
        key_index = head.index(key_field)
        deposit_key = next((n for n in ("Majority protein IDs", "Protein IDs", "Gene names")
                            if n in group_column), None)
        lookup = {}
        for row in groups.rows:
            for part in row[group_column[deposit_key]].split(";"):
                lookup.setdefault(part.strip(), row)
        drawn_total = matched = 0
        for body_row in body[head_index + 1:]:
            key = str(body_row[key_index]).split(";")[0].strip() if body_row[key_index] else ""
            group = lookup.get(key)
            if group is None:
                continue
            matched += 1
            for name in lfq:
                deposit_name = name.replace("-", "_")
                text = group[group_column[deposit_name]].strip() if deposit_name in group_column else ""
                try:
                    value = float(text)
                except ValueError:
                    value = 0.0
                if value <= 0:
                    drawn_total += 1
        print(f"  matched {matched} of {len(body) - head_index - 1} published protein claims to the deposit")
        print(f"  cells with no deposited LFQ value (drawn): {drawn_total} of {matched * len(lfq)}")
        print("  compare: 95% of published SITE claims contain a drawn cell")
else:
    print("  MOESM4 or the deposited proteinGroups is not in the raw store")

print("\nKK — a bounded scan for the paper's 2,341 modified peptides")
peptide_field = next((n for n in table.header if n.lower().startswith("mod. peptide id")), None)
hits = []
for localisation in (0.0, 0.5, 0.6, 0.7, 0.75, 0.8, 0.9, 0.95, 1.0):
    line = f"  loc >= {localisation:<5}"
    for score in (0, 20, 30, 40, 50, 60):
        union: set[str] = set()
        for row in rows:
            if raw_value(row, "Localization prob") < localisation:
                continue
            if "Score" in column and raw_value(row, "Score") < score:
                continue
            if "Reverse" in column and row[column["Reverse"]].strip() == "+":
                continue
            if "Potential contaminant" in column and row[column["Potential contaminant"]].strip() == "+":
                continue
            union.update(p for p in row[column[peptide_field]].split(";") if p.strip())
        line += f" {len(union):>6}"
        if len(union) == 2341:
            hits.append((localisation, score))
    print(line + f"   (scores {0, 20, 30, 40, 50, 60})")
print(f"  exact hits on 2,341: {len(hits)} {hits if hits else ''}")
if len(hits) == 1:
    localisation, score = hits[0]
    print(f"  the single hit is localisation >= {localisation}, score >= {score}; defensible on its"
          f" face only if both are round values or match a stated default.")
elif len(hits) > 1:
    print("  several cells hit the target, so the scan is not informative.")
else:
    print("  no cell reaches 2,341: on this axis the published record is exhausted, evidenced rather")
    print("  than assumed.")
