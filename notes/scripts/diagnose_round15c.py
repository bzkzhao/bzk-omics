"""Round-15c descriptive, UNREGISTERED: extending the protein census to MOESM5's 324 claims.

Usage, from the repository root:
    .venv/bin/python notes/scripts/diagnose_round15c.py 2>&1 | tee notes/logs/round15c.txt

Round 15b established the protein census on MOESM4's 25 claims: 44% contain a drawn cell, against
95% at site level. The denominator is small. MOESM5 holds 324 more published protein claims, but
its values are log2 RATIOS despite an "LFQ intensity" header, so cells cannot be read directly.

With 15b's column mapping confirmed by the published values themselves, the ratios can be tested
against the deposited `proteinGroups`: if each published value equals log2 of the deposited LFQ
value minus a per-row or per-column reference, the reference is recoverable and the census extends
from 25 claims to 349.

The script tests three readings before using any of them, and runs the census only on a reading the
published numbers confirm. If none is confirmed, MOESM5 stays out and the census stands at 25.
"""

import math
import warnings

import numpy as np
from openpyxl import load_workbook

from bzk.adapters import maxquant
from bzk.provenance.raw_store import raw_root
from bzk.sources.pxd018299_h10 import HOME

warnings.filterwarnings("ignore")

MAPPING = {
    "KO_IFN_1": "LFQ intensity KO_INF_P_2hGradient1",
    "KO_IFN_2": "LFQ intensity KO_INF_P_2hGradient2",
    "KO_IFN_3": "LFQ intensity KO_INF_P_2hGradient3",
    "WT_IFN_1": "LFQ intensity WT_INF_P_2hGradient1",
    "WT_IFN_2": "LFQ intensity WT_INF_P_2hGradient2",
    "WT_IFN_3": "LFQ intensity WT_INF_P_2hGradient3",
}

moesm5 = next(iter(raw_root(HOME).glob("*/41416_2020_1167_MOESM5_ESM.xlsx")), None)
groups_path = next(iter(raw_root(HOME).glob("*/HAP1_USP18KO_proteinGroups.txt")), None)
if not moesm5 or not groups_path:
    raise SystemExit("MOESM5 or the deposited proteinGroups is not in the raw store")

book = load_workbook(moesm5, read_only=True, data_only=True)
body = [r for r in book.worksheets[0].iter_rows(values_only=True) if any(c is not None for c in r)]
head_index = next(i for i, r in enumerate(body) if any("LFQ intensity" in str(c) for c in r if c))
head = [str(c).strip() if c is not None else "" for c in body[head_index]]
published_rows = body[head_index + 1:]
print(f"MOESM5 header: {head}")
print(f"published claims: {len(published_rows)}")

groups = maxquant.read_table(groups_path)
group_column = {n: i for i, n in enumerate(groups.header)}
key_field = next((n for n in head if "protein id" in n.lower()), None)
print(f"identifier column: {key_field!r}")
if key_field is None:
    raise SystemExit("no protein identifier column in MOESM5 — cannot join")
key_index = head.index(key_field)
lookup: dict[str, list[str]] = {}
for row in groups.rows:
    for part in row[group_column["Protein IDs"]].split(";"):
        lookup.setdefault(part.strip(), row)

columns = [name for name in head if name.startswith("LFQ intensity")]
sample_of = {name: name.replace("LFQ intensity", "").replace("-", "_").strip() for name in columns}
print(f"published sample columns, normalised: {list(sample_of.values())}")


def deposited(group: list[str], sample: str) -> float:
    name = MAPPING.get(sample)
    text = group[group_column[name]].strip() if name and name in group_column else ""
    try:
        return float(text)
    except ValueError:
        return 0.0


matched = []
for row in published_rows:
    key = str(row[key_index]).split(";")[0].strip() if row[key_index] else ""
    group = lookup.get(key)
    if group is not None:
        matched.append((row, group))
print(f"matched {len(matched)} of {len(published_rows)} claims to a deposited group")

print("\ntesting three readings of MOESM5's values, on cells the deposit provides")
readings: dict[str, list[tuple[float, float]]] = {"raw log2": [], "minus row mean": [], "minus column median": []}
column_values: dict[str, list[float]] = {sample: [] for sample in sample_of.values()}
for row, group in matched:
    for name in columns:
        value = deposited(group, sample_of[name])
        if value > 0:
            column_values[sample_of[name]].append(math.log2(value))
column_median = {sample: float(np.median(v)) if v else math.nan for sample, v in column_values.items()}

for row, group in matched:
    present = [(name, math.log2(deposited(group, sample_of[name])))
               for name in columns if deposited(group, sample_of[name]) > 0]
    if not present:
        continue
    row_mean = float(np.mean([v for _, v in present]))
    for name, deposit_log2 in present:
        try:
            published = float(row[head.index(name)])
        except (TypeError, ValueError):
            continue
        readings["raw log2"].append((published, deposit_log2))
        readings["minus row mean"].append((published, deposit_log2 - row_mean))
        readings["minus column median"].append((published, deposit_log2 - column_median[sample_of[name]]))

confirmed = None
for label, pairs in readings.items():
    if not pairs:
        print(f"  {label:>22}: no comparable cells")
        continue
    published_values, candidate = np.array(pairs).T
    difference = published_values - candidate
    share = float((np.abs(difference) <= 0.01).mean())
    print(f"  {label:>22}: {len(pairs)} cells, within 0.01: {share:.1%}, "
          f"median difference {np.median(difference):+.3f}, correlation "
          f"{np.corrcoef(published_values, candidate)[0, 1]:.4f}")
    if share >= 0.95 and confirmed is None:
        confirmed = label

if confirmed is None:
    print("\n  No reading is confirmed by the published values. MOESM5 stays out of the census,")
    print("  which therefore stands at MOESM4's 25 claims. What MOESM5's numbers are remains open.")
    raise SystemExit(0)

print(f"\n  confirmed reading: {confirmed}")
print("\nthe protein-grain census, extended")
drawn_claims = whole_arm = drawn_cells = 0
for row, group in matched:
    drawn = [deposited(group, sample) <= 0 for sample in MAPPING]
    drawn_cells += sum(drawn)
    drawn_claims += any(drawn)
    whole_arm += all(drawn[:3]) or all(drawn[3:])
print(f"  published protein claims examined: {len(matched)}")
print(f"    containing at least one drawn cell: {drawn_claims} ({drawn_claims / len(matched):.0%})")
print(f"    with a whole arm drawn:             {whole_arm} ({whole_arm / len(matched):.0%})")
print(f"    drawn cells: {drawn_cells} of {len(matched) * 6} ({drawn_cells / (len(matched) * 6):.1%})")
print("  site level, for comparison: 95% of claims contain a drawn cell, 90% with a whole arm")
