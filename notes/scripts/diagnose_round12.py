"""Round-12 descriptive, UNREGISTERED checks (adversarial review, 2026-09-22).

Usage, from the repository root:  .venv/bin/python notes/scripts/diagnose_round12.py

Not an independent path where it re-runs the pipeline; nothing registered changes.

Y — THE TWO UNOPENED SUPPLEMENTS. The raw store holds `41416_2020_1167_MOESM4_ESM.xlsx` (16 kB) and
    `MOESM5_ESM.xlsx` (74 kB) for the anchor, and neither has been read in this project. If either
    carries a parameters sheet, a processed matrix or a second claim list, it bears directly on the
    threshold and the population, which are the two conditions the instability figure does not meet.

Z — WHY IS id 889 NEVER CALLED? Round 11 read the nine structural discordances. Eight sit at 2 to 4
    imputed values, which is draw territory. One does not: id 889 has SIX measured values, a
    localisation of 1 and a score of 89, and it is published yet called in no draw. A fully measured
    published claim that the reconstruction never calls is the only residual that no draw, threshold
    or population argument can absorb. This prints its numbers beside the published ones.
"""

import json
import math
import warnings

import numpy as np
from openpyxl import load_workbook

from bzk.adapters import maxquant
from bzk.provenance.raw_store import raw_root, verify
from bzk.sources.pxd018299_h10 import (
    HOME,
    PXD018299_SITES,
    SUPP_DATA_1,
    _float_cell,
    _s1_rows,
    _supplement_path,
    anchor_population,
    s1_columns,
)
from bzk.stats.imputation import downshifted_normal
from bzk.stats.perseus_s0 import _relabellings, _statistic

warnings.filterwarnings("ignore")
ALPHA, S0, R_PERM = 0.01, 0.1, 250
DEFAULT = {"downshift_sd": 1.8, "width_sd": 0.3, "scope": "per_sample"}
KO = ["Intensity KO_IFN_1", "Intensity KO_IFN_2", "Intensity KO_IFN_3"]
WT = ["Intensity WT_IFN_1", "Intensity WT_IFN_2", "Intensity WT_IFN_3"]

print("Y — the two unopened anchor supplements")
for name in ("41416_2020_1167_MOESM4_ESM.xlsx", "41416_2020_1167_MOESM5_ESM.xlsx"):
    matches = sorted(raw_root(HOME).glob(f"*/{name}"))
    if not matches:
        print(f"  {name}: not in the raw store")
        continue
    path = matches[0]
    book = load_workbook(path, read_only=True, data_only=True)
    print(f"  {name} ({path.stat().st_size:,} bytes), sheets: {book.sheetnames}")
    for sheet in book.worksheets:
        rows_seen = list(sheet.iter_rows(values_only=True))
        filled = [r for r in rows_seen if any(c is not None for c in r)]
        print(f"    sheet {sheet.title!r}: {len(filled)} non-empty rows")
        for row in filled[:4]:
            cells = [str(c)[:28] for c in row if c is not None][:9]
            print(f"      {cells}")

print("\nZ — id 889: measured in all six columns, published, never called")
table = maxquant.read_table(
    verify(PXD018299_SITES.expected_content_hash, filename=PXD018299_SITES.filename, home=HOME)
)
rows, column = anchor_population(table)
ids = [str(r[column["id"]]).strip() for r in rows]


def cell(row: list[str], name: str) -> float:
    raw = row[column[name]].strip() if name in column else ""
    try:
        value = float(raw)
    except ValueError:
        return math.nan
    return math.log2(value) if value > 0 else math.nan


matrix = np.array([[cell(r, n) for n in KO + WT] for r in rows])
measured = ~np.isnan(matrix)
keep = (measured[:, :3].sum(1) >= 1) | (measured[:, 3:].sum(1) >= 1)
kept = matrix[keep]
kept_ids = [i for i, k in zip(ids, keep, strict=True) if k]
position = kept_ids.index("889")

with open("tests/fixtures/pxd018299_published_cascade.json") as handle:
    cascade = json.load(handle)["rows"]
claim = next(c for c in cascade if str(c["deposit_id"]) == "889")
header, s1 = _s1_rows(_supplement_path(SUPP_DATA_1, HOME))
published_row = s1[int(claim["row"])]
published_values = [_float_cell(published_row.get(n)) for n in s1_columns(header)]
print(f"  deposit values (log2): {[round(v, 3) for v in kept[position]]}")
print(f"  S1 values:             {[round(v, 3) for v in published_values]}")
print(f"  difference:            {[round(a - b, 4) for a, b in zip(kept[position], published_values, strict=True)]}")
print(f"  KO mean {np.mean(kept[position][:3]):.3f}, WT mean {np.mean(kept[position][3:]):.3f}, "
      f"difference {np.mean(kept[position][:3]) - np.mean(kept[position][3:]):+.3f}")
print(f"  within-group SDs: KO {np.std(kept[position][:3], ddof=1):.3f}, "
      f"WT {np.std(kept[position][3:], ddof=1):.3f}")
for name in ("Localization prob", "Score", "Delta score", "PEP", "Proteins", "Gene names",
             "Position", "Multiplicity", "Reverse", "Potential contaminant"):
    if name in column:
        print(f"  deposit {name}: {rows[ids.index('889')][column[name]]!r}")

statistics, ranks_of = [], []
for seed in range(20):
    filled = downshifted_normal(kept, seed=seed, **DEFAULT).values
    observed = _statistic(filled[:, :3], filled[:, 3:], S0)
    finite = np.isfinite(observed)
    magnitude = np.abs(observed[finite])
    thresholds = np.sort(magnitude)
    counts = []
    for assignment in _relabellings(3, 3, scheme="random_excluding_trivial",
                                    randomisations=R_PERM, seed=seed)[0]:
        mask = np.zeros(6, dtype=bool)
        mask[np.asarray(assignment[:3])] = True
        null = np.abs(_statistic(filled[:, mask], filled[:, ~mask], S0))[finite]
        null = np.sort(null[np.isfinite(null)])
        counts.append(null.size - np.searchsorted(null, thresholds, side="left"))
    null_count = np.array(counts, dtype=float).mean(axis=0)
    observed_count = magnitude.size - np.searchsorted(thresholds, thresholds, side="left")
    q = np.full(observed.shape, np.inf)
    q[finite] = np.minimum.accumulate(null_count / np.maximum(observed_count, 1))[
        np.searchsorted(thresholds, magnitude, side="right") - 1
    ]
    statistics.append((observed[position], q[position]))
    ranks_of.append(int(np.argsort(np.argsort(-np.where(observed > 0, observed, -np.inf)))[position]))
d_values = [d for d, _ in statistics]
q_values = [q for _, q in statistics]
print(f"  statistic across 20 draws: {min(d_values):+.3f} to {max(d_values):+.3f} "
      f"(fully measured, so it should not move at all)")
print(f"  q across 20 draws: {min(q_values):.4f} to {max(q_values):.4f}  (threshold {ALPHA})")
print(f"  positive-direction rank: {min(ranks_of)} to {max(ranks_of)} of {int(keep.sum())} rows")
print("  A fully measured claim whose statistic is stable but whose q sits above the threshold is")
print("  a selection difference, not a draw effect: the paper called it and this reconstruction")
print("  does not, at any draw.")
