"""`bzk/pxd055843_depletion_ranking.py` — the ordering, checked for internal relations only.

**The artefact is in no content store in this container and no `.xlsx` is in the tree**, so the
checks here split the way `tests/test_pxd055843_perseus.py` splits: the mechanics run over
**synthetic** workbooks built in `tmp_path` with the deposit's recorded shape — three header rows,
a merged qualifier over the quantitative columns, stamped statistics and identifier columns — and
the artefact path skips when the bytes are absent, exactly as
`tests/test_pxd018299_baseline.py` does:

    try:
        path = deposit_path()
    except (FileNotFoundError, ValueError) as exc:
        pytest.skip(f"deposit not in raw/ — run `python -m bzk.sources.pride` first ({exc})")

One clause is added to that and the reason is stated: this module locates its bytes through
`bzk/sources/pxd055843_perseus.py`'s `locate`, which converts the store's `FileNotFoundError` into
a `SystemExit` carrying the digest it looked for. Catching only the two the precedent names would
turn an absent artefact into a red test on every fresh session, which is the failure that docstring
says the skip exists to prevent.

**Nothing here pins a gene, a rank, a difference or a q-value of the artefact.** Which protein is
most depleted in this contrast is the finding the module exists to compute, and a test that
asserted it would be the answer written down before anyone read the output — which is also why the
module commits no fixture. The synthetic files' numbers are this file's own and are asserted as
*relations* (sorted, excluded, agrees with the lookup), never as the artefact's values.

**The shape being modelled is reviewer-supplied and not re-derivable in this container**, and is
labelled so wherever it is stated.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest

from bzk.pxd055843_depletion_ranking import (
    DIFFERENCE_MARK,
    GENE_MARK,
    HEAD_TAIL,
    MINUS_LOG_P_MARK,
    Q_VALUE_MARK,
    Ranked,
    RankingError,
    column,
    header_row,
    lookup,
    members,
    rank,
    report,
)
from bzk.sources.pxd055843_perseus import locate

SUFFIX = "siUSP24_IFN_siCTRL_IFN"

#: The synthetic file's data rows: gene, difference, q-value, -log p. Deliberately out of order,
#: with one null difference spelled `NaN` and one spelled as an empty cell, and one identifier cell
#: carrying two members. The values are this file's own and say nothing about the artefact.
DATA: list[tuple[str, Any, Any, Any]] = [
    ("AAA", 1.25, 0.04, 1.40),
    ("BBB", -3.10, 0.001, 3.00),
    ("CCC;DDD", -0.50, 0.22, 0.66),
    ("EEE", "NaN", 0.10, 1.00),
    ("FFF", 2.75, 0.002, 2.70),
    ("GGG", None, None, None),
    ("HHH", -3.10, 0.05, 1.30),
]


def _under(name: str, entry: tuple[str, Any, Any, Any]) -> Any:
    """The value a data row carries under one stamped column, chosen by the column's own name.

    **By name and not by position, for the same reason the module under test discovers columns that
    way.** A builder that laid the four values out in a fixed order put them under the wrong
    headings the moment a test added a fifth column — the `p-value` file below did exactly that,
    and its q-value column silently read empty while the test still passed on what it asserted.

    `MINUS_LOG_P_MARK` is tested first because it contains the plain p-value spelling; a column
    matching no mark carries nothing, which is what the plain `p-value` column is here for.
    """
    gene, difference, q_value, minus_log_p = entry
    for mark, value in (
        (MINUS_LOG_P_MARK, minus_log_p),
        (DIFFERENCE_MARK, difference),
        (Q_VALUE_MARK, q_value),
        (GENE_MARK, gene),
    ):
        if mark in name:
            return value
    return None


def _book(
    path: Path,
    *,
    stamped: list[str] | None = None,
    label_rows: int = 2,
    data: list[tuple[str, Any, Any, Any]] | None = None,
) -> Path:
    """A workbook of the deposit's shape. Reviewer-supplied shape; the values are this file's own.

    `label_rows` rows of merged group labels, then the stamped row, then the data. Two quantitative
    columns stand to the left of the stamped ones so that the statistics are not at index 0 and a
    reader taking a position rather than a name would be visibly wrong.
    """
    import openpyxl

    names = (
        stamped
        if stamped is not None
        else [
            f"T: {GENE_MARK}s",
            f"N: {DIFFERENCE_MARK} {SUFFIX}",
            f"N: {Q_VALUE_MARK} {SUFFIX}",
            f"N: {MINUS_LOG_P_MARK} {SUFFIX}",
        ]
    )
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    width = 2 + len(names)
    for index in range(label_rows):
        sheet.append([f"Set {index + 1}", None] + [None] * len(names))
    sheet.append(["/raw/A.d", "/raw/B.d", *names])
    for entry in data if data is not None else DATA:
        sheet.append([100.0, 200.0, *(_under(name, entry) for name in names)])
    if label_rows:
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    assert sheet.max_column == width, (
        "the builder laid out a row of a different width than its header"
    )
    workbook.save(path)
    return path


# --------------------------------------------------------------------------------------------
# Column discovery
# --------------------------------------------------------------------------------------------


def test_each_needed_mark_matches_exactly_one_column(tmp_path: Path) -> None:
    """The four marks resolve, and the matched name is reported rather than only the index.

    `columns` is asserted because the printed report shows it: an operator reading the output has
    to be able to see *which* header the module read, not merely that it read one.
    """
    ranked = rank(_book(tmp_path / "export.xlsx"))
    assert set(ranked.columns) == {GENE_MARK, DIFFERENCE_MARK, Q_VALUE_MARK, MINUS_LOG_P_MARK}
    assert ranked.columns[DIFFERENCE_MARK] == f"{DIFFERENCE_MARK} {SUFFIX}"
    assert ranked.columns[GENE_MARK] == f"{GENE_MARK}s"


def test_a_mark_matching_nothing_is_refused_and_the_header_is_named(tmp_path: Path) -> None:
    """An absent column fails loudly. The header is in the message because it is what to look at."""
    book = _book(
        tmp_path / "no-gene.xlsx",
        stamped=[
            f"N: {DIFFERENCE_MARK} {SUFFIX}",
            f"N: {Q_VALUE_MARK} {SUFFIX}",
            f"N: {MINUS_LOG_P_MARK} {SUFFIX}",
            "T: Protein.Group",
        ],
    )
    with pytest.raises(RankingError) as exc:
        rank(book)
    assert GENE_MARK in str(exc.value)
    assert "Protein.Group" in str(exc.value)


def test_a_mark_matching_twice_is_refused_rather_than_taking_the_first(tmp_path: Path) -> None:
    """Two matches is the failure the rule exists for — see the module docstring on `p-value`.

    **The workbook carries every other needed column on purpose.** A first draft of this omitted
    the `-Log` one, so removing the several-match refusal still raised `RankingError` — for the
    *absent* column — and the message, which names the whole header, still contained both
    difference spellings. The test passed against a module that took the first match, which is
    `CLAUDE.md` point 2's vacuous pass. With every other mark resolvable, the duplicate is the only
    thing left that can fail, and `column` is also called directly on a two-match header, where no
    other failure is reachable at all.
    """
    book = _book(
        tmp_path / "twice.xlsx",
        stamped=[
            f"T: {GENE_MARK}s",
            f"N: {DIFFERENCE_MARK} {SUFFIX}",
            f"N: {DIFFERENCE_MARK} other_contrast",
            f"N: {Q_VALUE_MARK} {SUFFIX}",
            f"N: {MINUS_LOG_P_MARK} {SUFFIX}",
        ],
    )
    with pytest.raises(RankingError) as exc:
        rank(book)
    assert f"{DIFFERENCE_MARK} {SUFFIX}" in str(exc.value)
    assert f"{DIFFERENCE_MARK} other_contrast" in str(exc.value)
    with pytest.raises(RankingError):
        column([f"{DIFFERENCE_MARK} one", f"{DIFFERENCE_MARK} two"], DIFFERENCE_MARK)


def test_the_plain_p_value_column_does_not_collide_with_the_minus_log_one(tmp_path: Path) -> None:
    """The named collision, asserted rather than noted: both columns present, one match each.

    `Student's T-test p-value X` is a substring of `-Log Student's T-test p-value X`, so a mark
    without the `-Log` would match both. This is the file shape that would prove it.
    """
    book = _book(
        tmp_path / "both-p.xlsx",
        stamped=[
            f"T: {GENE_MARK}s",
            f"N: {DIFFERENCE_MARK} {SUFFIX}",
            f"N: Student's T-test p-value {SUFFIX}",
            f"N: {MINUS_LOG_P_MARK} {SUFFIX}",
            f"N: {Q_VALUE_MARK} {SUFFIX}",
        ],
    )
    ranked = rank(book)
    assert ranked.columns[MINUS_LOG_P_MARK] == f"{MINUS_LOG_P_MARK} {SUFFIX}"
    with pytest.raises(RankingError):
        column([f"Student's T-test p-value {SUFFIX}", f"{MINUS_LOG_P_MARK} {SUFFIX}"], "p-value")


def test_the_header_row_is_found_by_its_stamp_and_not_by_its_index(tmp_path: Path) -> None:
    """A file with a different number of label rows still reads, and one with none still fails.

    The deposit's export puts the stamped row at index 2; taking that index would be a coincidence
    of one file rather than a rule, so the rule is asserted by moving the row.
    """
    assert rank(_book(tmp_path / "one-label.xlsx", label_rows=1)).contrast == SUFFIX
    assert rank(_book(tmp_path / "no-label.xlsx", label_rows=0)).contrast == SUFFIX
    assert header_row([["Set 1", ""], ["a", "b"], ["T: Genes", "N: x"]]) == 2
    with pytest.raises(RankingError):
        header_row([["Set 1", ""], ["a", "b"], ["Genes", "x"]])


def test_the_contrast_is_read_off_the_file_and_not_restated(tmp_path: Path) -> None:
    """A file spelling a different suffix reports that suffix. Nothing here holds a constant."""
    other = _book(
        tmp_path / "other.xlsx",
        stamped=[
            f"T: {GENE_MARK}s",
            f"N: {DIFFERENCE_MARK} a_vs_b",
            f"N: {Q_VALUE_MARK} a_vs_b",
            f"N: {MINUS_LOG_P_MARK} a_vs_b",
        ],
    )
    assert rank(other).contrast == "a_vs_b"


# --------------------------------------------------------------------------------------------
# The ordering
# --------------------------------------------------------------------------------------------


def test_the_ranking_is_ascending_by_difference_with_ranks_running_from_one(tmp_path: Path) -> None:
    """Ascending, so rank 1 is the most negative — and the ranks are the positions, not a field."""
    ranked = rank(_book(tmp_path / "export.xlsx"))
    assert ranked.rows, "an empty ordering passes both checks below vacuously"
    differences = [row.difference for row in ranked.rows]
    assert differences == sorted(differences)
    assert [row.rank for row in ranked.rows] == list(range(1, len(ranked.rows) + 1))


def test_rows_with_no_difference_are_excluded_and_counted(tmp_path: Path) -> None:
    """A `NaN` and an empty cell are both absences, and neither sorts. See the module docstring.

    The counts are asserted as a partition of the data rows rather than as three numbers, so the
    check cannot pass with a row that fell out of both.
    """
    ranked = rank(_book(tmp_path / "export.xlsx"))
    assert ranked.data_rows == len(DATA)
    assert len(ranked.rows) + ranked.without_difference == ranked.data_rows
    assert ranked.without_difference == sum(1 for row in DATA if row[1] in (None, "NaN"))
    assert not [row for row in ranked.rows if row.gene in ("EEE", "GGG")]


def test_a_cell_that_is_neither_empty_nor_a_number_is_refused(tmp_path: Path) -> None:
    """Refused rather than read as absent — a value nobody can parse is not a value that is gone."""
    book = _book(tmp_path / "junk.xlsx", data=[("AAA", "filtered", 0.04, 1.4)])
    with pytest.raises(RankingError) as exc:
        rank(book)
    assert "filtered" in str(exc.value)


def test_ties_keep_the_file_s_row_order(tmp_path: Path) -> None:
    """The stable sort the module depends on, asserted because nothing else enforces it."""
    ranked = rank(_book(tmp_path / "export.xlsx"))
    tied = [row.gene for row in ranked.rows if row.difference == -3.10]
    assert tied == ["BBB", "HHH"]


def test_a_missing_q_value_is_the_third_state_and_not_a_zero(tmp_path: Path) -> None:
    """`None` on the row and `absent` in the report — a zero would be a q-value nobody measured."""
    book = _book(tmp_path / "no-q.xlsx", data=[("AAA", -1.0, None, None)])
    ranked = rank(book)
    assert ranked.rows[0].q_value is None
    assert ranked.rows[0].minus_log_p is None
    assert any("absent" in line for line in report(ranked))


# --------------------------------------------------------------------------------------------
# The gene lookup
# --------------------------------------------------------------------------------------------


def _agrees(ranked: Ranked, wanted: int) -> None:
    """The row at `wanted` is the row the lookup returns for its own gene, at that same rank."""
    row = ranked.rows[wanted - 1]
    assert row.rank == wanted
    hits = lookup(ranked, members(row.gene)[0])
    assert row in hits
    assert all(ranked.rows[hit.rank - 1] == hit for hit in hits)


def test_the_lookup_returns_the_row_the_ranking_holds_at_that_rank(tmp_path: Path) -> None:
    """Every rank, not a chosen one: the lookup and the ordering are two readings of one list."""
    ranked = rank(_book(tmp_path / "export.xlsx"))
    assert ranked.rows
    for wanted in range(1, len(ranked.rows) + 1):
        _agrees(ranked, wanted)


def test_the_lookup_reads_a_cell_s_members_and_reports_every_match(tmp_path: Path) -> None:
    """A name sharing its cell is found, and a list is returned rather than a first hit."""
    ranked = rank(_book(tmp_path / "export.xlsx"))
    both = [lookup(ranked, name) for name in ("CCC", "DDD")]
    assert both[0] == both[1]
    assert len(both[0]) == 1
    assert lookup(ranked, "CCC;DDD") == []


def test_an_absent_gene_is_said_to_be_absent_rather_than_omitted(tmp_path: Path) -> None:
    """The third state reaches the report. A gene nobody finds is a line, not a silence."""
    ranked = rank(_book(tmp_path / "export.xlsx"))
    assert lookup(ranked, "ZZZ") == []
    assert any("ZZZ: absent" in line for line in report(ranked, ["ZZZ"]))


def test_the_lookup_is_case_sensitive_as_the_report_says_it_is(tmp_path: Path) -> None:
    """Asserted because the report claims it in words, and a claim in words is not a check."""
    ranked = rank(_book(tmp_path / "export.xlsx"))
    assert lookup(ranked, "aaa") == []
    assert len(lookup(ranked, "AAA")) == 1


def test_the_report_says_when_the_two_blocks_overlap(tmp_path: Path) -> None:
    """Fewer ranked rows than two blocks means the same ordering is printed twice; it says so."""
    ranked = rank(_book(tmp_path / "export.xlsx"))
    assert len(ranked.rows) <= 2 * HEAD_TAIL
    assert any("overlap" in line for line in report(ranked))


# --------------------------------------------------------------------------------------------
# The artefact: needs the bytes in raw/
# --------------------------------------------------------------------------------------------


@pytest.fixture(scope="module")
def artefact() -> Ranked:
    try:
        deposit = locate()
    except (SystemExit, FileNotFoundError, ValueError) as exc:
        pytest.skip(f"the artefact is not in raw/ — it has to be stored by hand ({exc})")
    return rank(deposit)


def test_the_artefact_ranks_and_its_counts_partition_its_rows(artefact: Ranked) -> None:
    """The only non-synthetic check, and it pins nothing — see the module docstring.

    Sortedness, a rank sequence, a partition of the data rows and a contrast the file names. Every
    one is a relation the module must hold for any input; none is a value of this file.
    """
    assert artefact.rows, "an empty ordering passes the checks below vacuously"
    differences = [row.difference for row in artefact.rows]
    assert differences == sorted(differences)
    assert [row.rank for row in artefact.rows] == list(range(1, len(artefact.rows) + 1))
    assert len(artefact.rows) + artefact.without_difference == artefact.data_rows
    assert artefact.contrast
    assert artefact.columns[DIFFERENCE_MARK].startswith(DIFFERENCE_MARK)


def test_the_artefact_s_lookup_agrees_with_its_ordering_at_both_ends(artefact: Ranked) -> None:
    """The ends and the middle, taken *from* the ordering rather than named here."""
    assert artefact.rows
    for wanted in (1, (len(artefact.rows) + 1) // 2, len(artefact.rows)):
        _agrees(artefact, wanted)
