"""`bzk/sources/pxd026748_h5c_h9p.py` — H5c and H9p, on synthetic inputs.

**Entirely offline, and nothing here touches PXD026748.** Both hypotheses read a committed fixture,
a pinned supplement and two deposits that exist only on bzk's machine; what can be tested here is
the arithmetic between reading them and the verdict — which population is counted, which comparison
is made, what happens at a registered boundary, and what the two verdicts refuse to read.

**Every case is built to discriminate rather than to pass.** The cluster-exclusion case is one
where including cluster 3 would give a *different* verdict, not merely a different number; the
conditionality case is one where the unconditional comparison inverts the finding. A case that
agrees under both readings establishes nothing about which reading the code applies.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import numpy as np
import pytest
from openpyxl import Workbook

from bzk.sources import pxd026748_h5c_h9p as hyp
from bzk.sources.pxd026748_reconstruction import SampleAxes

GENOTYPES = ("WT", "ISG15-/-")
TREATMENTS = ("plpro_wt", "plpro_mut")

#: A synthetic design, only ever used where a `SampleAxes` is structurally required.
AXES = SampleAxes(
    labels=tuple(f"Intensity {i}" for i in range(12)),
    genotype=tuple(GENOTYPES[0] if i < 6 else GENOTYPES[1] for i in range(12)),
    treatment=tuple(TREATMENTS[i // 3 % 2] for i in range(12)),
)

#: Four members, not 360: the categories depend on the member count and on nothing else, so a
#: smaller family makes the same distinctions in a tenth of the lines.
MEMBERS = 4


def _claim(deposit_id: str, row: int, cluster: str, support: int) -> dict[str, Any]:
    """One committed claim, with `support` of `MEMBERS` members below the primary threshold."""
    return {
        "deposit_id": deposit_id,
        "published_row": row,
        "cluster": cluster,
        "uniprot_id": f"SYNTHETIC-{row:03d}",
        "min_p": [0.001] * support + [0.5] * (MEMBERS - support),
    }


def _flags(rows: dict[int, bool]) -> hyp.PublishedFlags:
    return hyp.PublishedFlags(
        column="In vivo ISG15 targets",
        flagged=dict(rows),
        values_as_found={"": sum(1 for v in rows.values() if not v), "x": sum(rows.values())},
    )


def _block(
    claims: list[dict[str, Any]], flagged_rows: dict[int, bool], *, missing: list[int] | None = None
) -> dict[str, Any]:
    """`h5c_block` over synthetic claims, with a GG matrix carrying the given missing counts."""
    values = np.zeros((len(claims), 12))
    for i, count in enumerate(missing or [0] * len(claims)):
        values[i, :count] = np.nan
    return hyp.h5c_block(
        claims=claims,
        categories=hyp.recomputed_categories(claims, member_count=MEMBERS),
        flags=_flags(flagged_rows),
        values=values,
        axes=AXES,
        row_of_id={c["deposit_id"]: i for i, c in enumerate(claims)},
    )


def _population(
    *, flagged_durable: int, flagged_total: int, unflagged_durable: int, unflagged_total: int
) -> tuple[list[dict[str, Any]], dict[int, bool]]:
    """An ISG15-cluster population with the given durable counts in each group."""
    claims: list[dict[str, Any]] = []
    flags: dict[int, bool] = {}
    row = 1
    for flagged, durable, total in (
        (True, flagged_durable, flagged_total),
        (False, unflagged_durable, unflagged_total),
    ):
        for i in range(total):
            claims.append(_claim(str(row), row, "Cluster 1a", MEMBERS if i < durable else 1))
            flags[row] = flagged
            row += 1
    return claims, flags


# ── H5c · the population ────────────────────────────────────────────────────────────────────────


def test_h5c_excludes_cluster_3(*, _unused: None = None) -> None:
    """*"Cluster 3's ubiquitin sites are excluded, because the flag concerns ISG15 targets"* — §5.

    The case is built so that including them **changes the verdict** and not merely a proportion:
    on clusters 1a/1b/2 alone the flagged group is 5/10 durable against 4/10, a difference of
    +0.10 and `discriminates`; ten flagged cluster-3 claims, none durable, take the flagged
    proportion to 5/20 and the difference to −0.15, which is `inverted`.
    """
    claims, flags = _population(
        flagged_durable=5, flagged_total=10, unflagged_durable=4, unflagged_total=10
    )
    for i in range(10):
        row = 100 + i
        claims.append(_claim(str(row), row, "Cluster 3", 1))
        flags[row] = True

    block = _block(claims, flags)

    assert block["population"]["claims"] == 20
    assert block["population"]["excluded_claims"] == 10
    assert block["group_sizes"] == {"flagged": 10, "unflagged": 10}
    assert block["verdict"] == "discriminates"


def test_h5c_counts_every_isg15_cluster() -> None:
    """1a, 1b and 2 are all in the population — an exclusion that took only 1a would also pass the
    test above, so the three are asserted by name."""
    claims = [
        _claim("1", 1, "Cluster 1a", MEMBERS),
        _claim("2", 2, "Cluster 1b", MEMBERS),
        _claim("3", 3, "Cluster 2", 0),
        _claim("4", 4, "Cluster 3", MEMBERS),
    ]
    block = _block(claims, {1: True, 2: False, 3: False, 4: True})

    assert block["population"]["by_cluster"] == {"Cluster 1a": 1, "Cluster 1b": 1, "Cluster 2": 1}


def test_h5c_stops_on_a_claim_that_joins_no_published_row() -> None:
    """The join is on `#`, and a claim with no published row has no group."""
    claims = [_claim("1", 1, "Cluster 1a", MEMBERS), _claim("2", 77, "Cluster 2", 0)]

    with pytest.raises(hyp.HypothesisError, match="published row 77"):
        _block(claims, {1: True})


# ── H5c · the verdict ───────────────────────────────────────────────────────────────────────────


def test_h5c_verdict_at_the_registered_boundaries() -> None:
    """±0.10 are inside the outer verdicts and +0.09 is not — §5's *"at least"* and *"at most"*.

    `0.5 - 0.4` rather than the literal `0.10`: that subtraction is `0.09999999999999998` in binary
    floating point, and it is the shape a real boundary case arrives in. A verdict that read the
    literal alone would call the registered case `does_not_discriminate` on a 2e-17 shortfall.
    """
    assert hyp.h5c_verdict(0.10) == "discriminates"
    assert hyp.h5c_verdict(0.5 - 0.4) == "discriminates"
    assert hyp.h5c_verdict(-0.10) == "inverted"
    assert hyp.h5c_verdict(-(0.5 - 0.4)) == "inverted"
    assert hyp.h5c_verdict(0.09) == "does_not_discriminate"
    assert hyp.h5c_verdict(-0.09) == "does_not_discriminate"
    assert hyp.h5c_verdict(0.0) == "does_not_discriminate"


def test_h5c_reads_durability_from_the_fixture_and_never_runs_a_member(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """§5: the durable proportion is *"taken from the committed reconstruction fixture, which is
    not re-run"*.

    A rerun would draw new values from the same grid and give slightly different categories, so it
    could not corroborate the fixture — it would replace it inside a hypothesis about it. The spy
    is on the module's own `run_member` name, which is what a call from `h5c_block` would reach.
    """
    calls: list[object] = []

    def _spy(*args: object, **kwargs: object) -> np.ndarray:
        calls.append(args)
        return np.zeros(1)

    monkeypatch.setattr(hyp, "run_member", _spy)
    claims, flags = _population(
        flagged_durable=5, flagged_total=10, unflagged_durable=4, unflagged_total=10
    )

    block = _block(claims, flags)

    assert calls == []
    assert block["verdict"] == "discriminates"


def test_h5c_confound_is_reported_and_never_read_by_the_verdict() -> None:
    """The declared confound is descriptive — §5 declares it and declares it descriptive at once.

    Two populations with identical durability and opposite missingness give the same difference and
    the same verdict, and differ only in the block the verdict does not read.
    """
    claims, flags = _population(
        flagged_durable=5, flagged_total=10, unflagged_durable=4, unflagged_total=10
    )
    none_missing = _block(claims, flags, missing=[0] * 20)
    many_missing = _block(claims, flags, missing=[9] * 10 + [0] * 10)

    assert none_missing["difference"] == many_missing["difference"]
    assert none_missing["verdict"] == many_missing["verdict"] == "discriminates"
    assert none_missing["confound_missing_values"] != many_missing["confound_missing_values"]
    assert many_missing["confound_missing_values"]["by_group"]["flagged"] == {"9": 10}


# ── H5c · the published flag ────────────────────────────────────────────────────────────────────


def _table_1(path: Path, rows: list[list[Any]]) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "Table 1"
    sheet.append(["Supplementary Table 1 — synthetic"])
    sheet.append(
        [
            "#",
            "Cluster",
            "Uniprot ID",
            "Lysine position",
            "Ubiquitin sites",
            "In vivo ISG15 targets",
        ]
    )
    for row in rows:
        sheet.append(row)
    book.save(path)
    return path


def test_the_flag_is_x_and_every_distinct_value_is_reported(tmp_path: Path) -> None:
    """`x` flags, and nothing else does — not `-`, not a blank, not another mark. Every value the
    column held is counted, so the stripping and case-folding this module chose are auditable."""
    path = _table_1(
        tmp_path / "t1.xlsx",
        [
            [1, "Cluster 1a", "SYNTHETIC-1", 10, "x", "x"],
            [2, "Cluster 1a", "SYNTHETIC-2", 11, "", " X "],
            [3, "Cluster 2", "SYNTHETIC-3", 12, "", None],
            [4, "Cluster 2", "SYNTHETIC-4", 13, "", "-"],
        ],
    )

    flags = hyp.published_flags(path)

    assert flags.column == "In vivo ISG15 targets"
    assert flags.flagged == {1: True, 2: True, 3: False, 4: False}
    assert flags.values_as_found == {"": 1, "-": 1, " X ": 1, "x": 1}


def test_a_repeated_published_row_number_stops_the_run(tmp_path: Path) -> None:
    path = _table_1(
        tmp_path / "t1.xlsx",
        [
            [1, "Cluster 1a", "SYNTHETIC-1", 10, "", "x"],
            [1, "Cluster 2", "SYNTHETIC-2", 11, "", ""],
        ],
    )

    with pytest.raises(hyp.HypothesisError, match="row number 1 more than once"):
        hyp.published_flags(path)


# ── the fixture-consistency guard ───────────────────────────────────────────────────────────────


def _fixture(claims: list[dict[str, Any]], counts: dict[str, int]) -> dict[str, Any]:
    return {
        "gg_run": True,
        "family": {
            "members": [{"seed": i} for i in range(MEMBERS)],
            "claims": claims,
            "readouts": {"primary": {"counts": counts}},
        },
    }


def test_a_fixture_that_disagrees_with_its_own_counts_stops_the_run() -> None:
    """H5c is a statement about the committed fixture's contents; if the file's stored counts and
    its own per-claim numbers have parted company, which of the two H5c is about is undecided."""
    claims = [_claim("1", 1, "Cluster 1a", MEMBERS), _claim("2", 2, "Cluster 1a", 1)]
    honest = _fixture(claims, {"durable": 1, "underdetermined": 1, "unsupported": 0})
    lying = _fixture(claims, {"durable": 2, "underdetermined": 0, "unsupported": 0})

    assert hyp.fixture_consistency(honest)["agrees"] is True
    with pytest.raises(hyp.HypothesisError, match="disagrees with itself"):
        hyp.fixture_consistency(lying)


def test_a_claim_whose_min_p_list_is_the_wrong_length_is_refused() -> None:
    """A support count over the wrong denominator is not a support count, and a short list would
    otherwise read as a claim supported by fewer members than the family has."""
    claims = [_claim("1", 1, "Cluster 1a", MEMBERS)]
    claims[0]["min_p"] = claims[0]["min_p"][:-1]

    with pytest.raises(hyp.HypothesisError, match="against 4 member"):
        hyp.recomputed_categories(claims, member_count=MEMBERS)


# ── H9p ─────────────────────────────────────────────────────────────────────────────────────────


def _rows(
    total: int, with_missing: int, underdetermined_among_missing: int
) -> list[dict[str, Any]]:
    """`total` rows, `with_missing` of them carrying a missing value, of which some are
    underdetermined. Every complete row is durable, which is the protein arm's real shape."""
    rows: list[dict[str, Any]] = []
    for i in range(with_missing):
        underdetermined = i < underdetermined_among_missing
        rows.append(
            {
                "accession": f"SYNTHETIC-{i:03d}",
                "support": 2 if underdetermined else MEMBERS,
                "category": "underdetermined" if underdetermined else "durable",
                "missing_values": 3,
            }
        )
    for i in range(total - with_missing):
        rows.append(
            {
                "accession": f"SYNTHETIC-C{i:03d}",
                "support": MEMBERS,
                "category": "durable",
                "missing_values": 0,
            }
        )
    return rows


def test_h9p_compares_conditional_on_missingness_not_unconditionally() -> None:
    """§5: *"the underdetermined proportion among claims whose row carries at least one missing
    value"*, because the protein arm's complete rows would otherwise settle the comparison.

    The case inverts between the two readings. Conditional: U_p = 3/4 = 0.75, U_s = 8/10 = 0.80,
    and 0.75 ≥ 0.8 × 0.80 = 0.64, so **general to the imputation step**. Unconditional: 3/20 = 0.15
    against 8/10 = 0.80, and 0.15 ≤ 0.5 × 0.80 = 0.40, which would read **particular to site
    data** — the opposite finding, from the same rows.
    """
    proteins = _rows(20, 4, 3)
    sites = _rows(10, 10, 8)

    block = hyp.h9p_block(proteins=proteins, sites=sites, join={})

    assert block["conditional"]["u_p"] == {"numerator": 3, "denominator": 4, "share": 0.75}
    assert block["conditional"]["u_s"] == {"numerator": 8, "denominator": 10, "share": 0.8}
    assert block["verdict"]["outcome"] == "general_to_imputation"
    # The unconditional pair is reported beside it and read by nothing.
    assert block["descriptive"]["unconditional_underdetermined"]["protein"]["share"] == 0.15
    assert block["descriptive"]["unconditional_underdetermined"]["site"]["share"] == 0.8
    assert block["descriptive"]["complete_rows"] == {"protein": 16, "site": 0}


def test_h9p_verdict_at_the_registered_factors() -> None:
    """`U_p <= 0.5 U_s` and `U_p >= 0.8 U_s`, both inclusive, with the band between them
    `indeterminate`."""
    u_s = {"numerator": 8, "denominator": 10, "share": 0.8}

    def outcome(share: float) -> str:
        verdict = hyp.h9p_verdict({"numerator": 1, "denominator": 2, "share": share}, u_s)
        return str(verdict["outcome"])

    assert outcome(0.40) == "particular_to_site_data"
    assert outcome(0.39) == "particular_to_site_data"
    assert outcome(0.64) == "general_to_imputation"
    assert outcome(0.65) == "general_to_imputation"
    assert outcome(0.50) == "indeterminate"


def test_h9p_is_undefined_rather_than_a_division_when_a_grain_has_no_missing_rows() -> None:
    """A proportion over no claims is not a small proportion. The reason names which side."""
    proteins = _rows(20, 4, 3)
    sites = _rows(10, 0, 0)

    block = hyp.h9p_block(proteins=proteins, sites=sites, join={})

    assert block["verdict"]["outcome"] == "undefined"
    assert "U_s" in str(block["verdict"]["reason"])
    assert block["conditional"]["u_s"] == {"numerator": 0, "denominator": 0, "share": None}


def test_the_h9p_verdict_reads_the_conditional_pair_and_nothing_else() -> None:
    """Two populations with the same conditional pair and opposite unconditional ones give the same
    verdict — the converse of the test above, and what makes "reported beside it" checkable."""
    sites = _rows(10, 10, 8)
    few_complete = hyp.h9p_block(proteins=_rows(5, 4, 3), sites=sites, join={})
    many_complete = hyp.h9p_block(proteins=_rows(400, 4, 3), sites=sites, join={})

    assert few_complete["verdict"] == many_complete["verdict"]
    assert (
        few_complete["descriptive"]["unconditional_underdetermined"]["protein"]["share"]
        != many_complete["descriptive"]["unconditional_underdetermined"]["protein"]["share"]
    )


# ── the fixture's shape ─────────────────────────────────────────────────────────────────────────


def test_the_recomputed_categories_are_the_registered_definition() -> None:
    """Durable is supported in **all** members, unsupported in none, underdetermined between."""
    claims = [
        _claim("1", 1, "Cluster 1a", MEMBERS),
        _claim("2", 2, "Cluster 1a", 0),
        _claim("3", 3, "Cluster 1a", 1),
        _claim("4", 4, "Cluster 1a", MEMBERS - 1),
    ]

    assert [e["category"] for e in hyp.recomputed_categories(claims, member_count=MEMBERS)] == [
        "durable",
        "unsupported",
        "underdetermined",
        "underdetermined",
    ]


def test_a_claim_with_no_row_in_the_population_is_unsupported_and_its_missingness_is_absent() -> (
    None
):
    """The reconstruction writes `null` min(P) for such a claim. A null clears no threshold, and
    its missing count is `absent` rather than zero — the two are not the same statement."""
    claims = [_claim("1", 1, "Cluster 1a", MEMBERS), _claim("2", 2, "Cluster 1a", 0)]
    claims[1]["min_p"] = [None] * MEMBERS
    values = np.zeros((1, 12))

    categories = hyp.recomputed_categories(claims, member_count=MEMBERS)
    counts = hyp.missing_counts(claims, values, {"1": 0})

    assert categories[1] == {"support": 0, "category": "unsupported"}
    assert counts == [0, None]


def test_main_writes_h5c_and_h9p_and_json_can_hold_them(tmp_path: Path) -> None:
    """The two blocks serialise. Asserted by writing and reading back, for the reason turn 15b's
    coerced-cell defect gives: a serialisation failure is invisible to an in-memory check."""
    claims, flags = _population(
        flagged_durable=5, flagged_total=10, unflagged_durable=4, unflagged_total=10
    )
    block = _block(claims, flags)
    h9p = hyp.h9p_block(proteins=_rows(20, 4, 3), sites=_rows(10, 10, 8), join={"exposure": 20})

    path = tmp_path / "out.json"
    path.write_text(json.dumps({"h5c": block, "h9p": h9p}, indent=2))
    read = json.loads(path.read_text())

    assert read["h5c"]["verdict"] == "discriminates"
    assert read["h9p"]["verdict"]["outcome"] == "general_to_imputation"
    assert len(read["h5c"]["claims"]) == 20
    assert len(read["h9p"]["proteins"]) == 20
    # **The 360 stored min(P) values are not copied through.** A committed claim carries them, and
    # `{**dict(claim)}` would duplicate the reconstruction fixture's whole matrix inside this one —
    # 3.5 MB of numbers this hypothesis recomputed into a support count and then does not use. The
    # count is kept; the matrix stays in the file whose hash this fixture pins.
    assert set(read["h5c"]["claims"][0]) == {
        "deposit_id",
        "published_row",
        "cluster",
        "uniprot_id",
        "lysine_position",
        "support",
        "category",
        "flagged",
        "missing_values",
    }


# ── end to end ──────────────────────────────────────────────────────────────────────────────────


def _workbook(path: Path, *, table_1: list[list[Any]], table_2_ids: list[str]) -> Path:
    """A supplement with the two sheets this module reads, each under a title row."""
    book = Workbook()
    sheet_1 = book.active
    sheet_1.title = "Table 1"
    sheet_1.append(["Supplementary Table 1 — synthetic"])
    sheet_1.append(
        [
            "#",
            "Cluster",
            "Uniprot ID",
            "Lysine position",
            "Ubiquitin sites",
            "In vivo ISG15 targets",
        ]
    )
    for row in table_1:
        sheet_1.append(row)

    sheet_2 = book.create_sheet("Table 2")
    sheet_2.append(["Supplementary Table 2 — synthetic"])
    sheet_2.append(["Uniprot ID", "Gene name"])
    for accession in table_2_ids:
        sheet_2.append([accession, "SYNTHETIC"])

    book.save(path)
    return path


def test_main_runs_both_hypotheses_end_to_end(tmp_path: Path) -> None:
    """`main()` over synthetic bytes: two deposits, two records, a supplement and a fixture.

    **This is the only test that executes `main`, `gg_matrix` and `shotgun_family`**, and that is
    its whole job — the arithmetic is established by the pure tests above. What it establishes is
    that the IO seam holds: the two arms are located and parsed, the two sheets are read, the
    committed fixture's claims join the published rows, the family runs over the shotgun matrix,
    and the result serialises. bzk runs the instrument; a signature error here would only surface
    there.

    The synthetic deposits and records come from `tests/test_pxd026748_reconstruction.py` rather
    than being rebuilt: they are the same two arms turn 16 built, and a second copy of them here
    would be two populations wearing one name.
    """
    # `test_pxd026748_reconstruction` by bare module name rather than `tests.test_...`: there is
    # no `tests/__init__.py`, so `mypy` maps this file to the bare name and rejects the dotted form
    # as a second module for one file. Adding the `__init__.py` would make `tests/` a package and
    # add a module to the tautology sweep's own count, which this turn is not to touch.
    from test_pxd026748_reconstruction import (
        CLAIM_IDS,
        _accession,
        _gg_bytes,
        _gg_rows,
        _record,
        _shotgun_bytes,
    )

    from bzk.provenance import raw_store
    from bzk.sources import protein_groups

    home = tmp_path / "home"
    fixtures_dir = tmp_path / "fixtures"
    fixtures_dir.mkdir()

    gg = raw_store.store(_gg_bytes(_gg_rows(1.0)), "GlyGly (K)Sites.txt", home=home)
    shotgun = raw_store.store(_shotgun_bytes(40), "proteinGroups.txt", home=home)
    gg_record = _record(
        tmp_path,
        name="curation_SYNTHETIC_gg.json",
        filename="GlyGly (K)Sites.txt",
        content_hash=gg.content_hash,
        prefix="Intensity ",
    )
    shotgun_record = _record(
        tmp_path,
        name="curation_SYNTHETIC_shotgun.json",
        filename="proteinGroups.txt",
        content_hash=shotgun.content_hash,
        prefix="LFQ intensity ",
    )

    # One published row per committed claim, the first two flagged, and one in cluster 3 so the
    # exclusion is exercised on the real path too.
    clusters = ["Cluster 1a", "Cluster 1b", "Cluster 2", "Cluster 2", "Cluster 3"]
    workbook = _workbook(
        tmp_path / "supp.xlsx",
        table_1=[
            [i + 1, clusters[i], f"SYNTHETIC-{i:03d}", 10 + i, "", "x" if i < 2 else ""]
            for i in range(len(CLAIM_IDS))
        ],
        table_2_ids=[_accession(i) for i in range(10)],
    )
    stored = raw_store.store(workbook.read_bytes(), "supp.xlsx", home=home)
    supplement = protein_groups.SupplementaryFile(
        label="synthetic",
        filename="supp.xlsx",
        expected_content_hash=stored.content_hash,
        doi="10.1038/s41590-021-01035-8",
    )

    claims = [
        _claim(row_id, i + 1, clusters[i], MEMBERS if i % 2 == 0 else 1)
        for i, row_id in enumerate(CLAIM_IDS)
    ]
    reconstruction = tmp_path / "reconstruction.json"
    reconstruction.write_text(
        json.dumps(
            _fixture(
                claims,
                {
                    "durable": sum(1 for i in range(len(claims)) if i % 2 == 0),
                    "underdetermined": sum(1 for i in range(len(claims)) if i % 2 == 1),
                    "unsupported": 0,
                },
            )
        ),
        encoding="utf-8",
    )

    code = hyp.main(
        home=home,
        fixtures_dir=fixtures_dir,
        digly_curation_path=gg_record,
        shotgun_curation_path=shotgun_record,
        reconstruction_fixture_path=reconstruction,
        supplement=supplement,
        progress=False,
    )
    written = json.loads((fixtures_dir / hyp.FIXTURE_NAME).read_text(encoding="utf-8"))

    assert code == 0
    assert written["generated_under"]["fixture_consistency"]["agrees"] is True
    # Four ISG15 claims of five, cluster 3 excluded on the real path.
    assert written["h5c"]["population"]["claims"] == 4
    assert written["h5c"]["group_sizes"] == {"flagged": 2, "unflagged": 2}
    assert written["h5c"]["verdict"] in {"discriminates", "does_not_discriminate", "inverted"}
    # Every claim carries the missing count the confound is reported from, read off the GG matrix.
    assert all(c["missing_values"] is not None for c in written["h5c"]["claims"])
    # H9p ran the registered family over the shotgun arm.
    assert written["h9p"]["join"]["members"] == 360
    assert written["h9p"]["join"]["table_2_ids_read"] == 10
    assert len(written["h9p"]["sites"]) == len(CLAIM_IDS)
    assert all(0 <= p["support"] <= 360 for p in written["h9p"]["proteins"])
